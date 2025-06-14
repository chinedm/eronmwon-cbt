import streamlit as st
import sqlite3
import pandas as pd
import datetime
import pdfplumber
from docx import Document
from striprtf.striprtf import rtf_to_text
from utils.db_access import connect_sqlite
from fpdf import FPDF

st.set_page_config(page_title="Data Center", layout="wide")
st.title("📃 Data Center: Manage, Import, Export Data")

conn = connect_sqlite()
cursor = conn.cursor()

if "logged_in" not in st.session_state or not st.session_state.logged_in:
    st.warning("Please log in to access this page.")
    st.stop()

menu = st.sidebar.radio("Choose Action", ["Capture Data", "Import Data", "Export Data", "Query Builder"])

# -----------------------------
# CAPTURE DATA SECTION
# -----------------------------
if menu == "Capture Data":
    table = st.selectbox("Select table to capture data into", ["Subjects", "Users"])

    # SUBJECTS FORM
    if table == "Subjects":
        st.subheader("📚 Add New Subject")
        subject_name = st.text_input("Subject Name")
        subject_desc = st.text_area("Description")
        if st.button("Add Subject"):
            if not subject_name:
                st.error("Subject Name is required.")
            else:
                cursor.execute("""
                    INSERT INTO subjects (name, description)
                    VALUES (?, ?)
                """, (subject_name, subject_desc))
                conn.commit()
                st.success(f"Subject '{subject_name}' added.")
                st.rerun()
        
        st.markdown("### View Subjects")
        cursor.execute("SELECT * FROM subjects ORDER BY id DESC LIMIT 10")
        subjects = cursor.fetchall()
        if subjects:
            for sub in subjects:
                st.markdown(f"**{sub['id']}** - {sub['name']}")
                st.markdown(f"Description: {sub['description']}")
        else:
            st.info("No subjects found. Please add a subject first.")

        st.markdown("### Manage Subjects")
        cursor.execute("SELECT * FROM subjects")
        for sub in cursor.fetchall():
            with st.expander(f"{sub['id']} - {sub['name']}"):
                new_name = st.text_input("Edit Name", sub['name'], key=f"en{sub['id']}")
                new_desc = st.text_area("Edit Description", sub['description'], key=f"ed{sub['id']}")
                c1, c2 = st.columns(2)
                if c1.button("Update", key=f"up{sub['id']}"):
                    cursor.execute("UPDATE subjects SET name=?, description=? WHERE id=?",
                                (new_name, new_desc, sub['id']))
                    conn.commit()
                    st.success(f"Subject '{new_name}' updated.")
                    st.rerun()
                if c2.button("Delete", key=f"del{sub['id']}"):
                    cursor.execute("DELETE FROM subjects WHERE id=?", (sub['id'],))
                    conn.commit()
                    st.success(f"Subject '{sub['name']}' deleted.")
                    st.rerun()

    # USERS FORM (unchanged, but make sure your users table matches this schema)
    elif table == "Users":
        st.subheader("👤 Register New User")
        user_id = st.text_input("User ID")
        first = st.text_input("First Name")
        last = st.text_input("Last Name")
        email = st.text_input("Email")
        phone = st.text_input("Phone")
        gender = st.selectbox("Gender", ["M", "F"])
        dob = st.date_input("DOB")
        institution = st.text_input("Institution")
        license_key = st.text_input("License Key")
        photo = st.text_input("Photo Path")
        if st.button("Register User"):
            if not user_id:
                st.error("User ID is required.")
            elif not first:
                st.error("First Name is required.")
            elif not last:
                st.error("Last Name is required.")
            else:
                cursor.execute("""
                    INSERT INTO users (user_id, first_name, last_name, email, phone_number, gender,
                    date_of_birth, photo, institution, license_key, registered_on)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (user_id, first, last, email, phone, gender, dob.isoformat(), photo, institution, license_key,
                      datetime.datetime.now().isoformat()))
                conn.commit()
                st.success("User registered.")
                st.rerun()

        st.markdown("### Manage Users")
        cursor.execute("SELECT * FROM users ORDER BY registered_on DESC LIMIT 10")
        for u in cursor.fetchall():
            with st.expander(f"{u['user_id']}: {u['first_name']} {u['last_name']}"):
                new_first = st.text_input("First Name", u['first_name'], key=f"ufn_{u['user_id']}")
                new_last = st.text_input("Last Name", u['last_name'], key=f"uln_{u['user_id']}")
                new_email = st.text_input("Email", u['email'], key=f"uem_{u['user_id']}")
                new_inst = st.text_input("Institution", u['institution'], key=f"uin_{u['user_id']}")
                new_license = st.text_input("License Key", u['license_key'], key=f"ulk_{u['user_id']}")
                col1, col2 = st.columns(2)
                if col1.button("Update", key=f"uu{u['user_id']}"):
                    cursor.execute("""
                        UPDATE users SET first_name=?, last_name=?, email=?, institution=?, license_key=?
                        WHERE user_id=?
                    """, (new_first, new_last, new_email, new_inst, new_license, u['user_id']))
                    conn.commit()
                    st.success(f"User '{new_first} {new_last}' updated.")
                    st.rerun()
                if col2.button("Delete", key=f"du{u['user_id']}"):
                    cursor.execute("DELETE FROM users WHERE user_id=?", (u['user_id'],))
                    conn.commit()
                    st.success(f"User '{u['first_name']} {u['last_name']}' deleted.")
                    st.rerun()

# -----------------------------
# IMPORT SECTION
# -----------------------------
elif menu == "Import Data":
    st.subheader("📄 Import from Excel, Word, PDF, etc.")
    file_type = st.selectbox("Select file type", ["Excel", "Word", "Text", "RTF", "PDF"])
    uploaded_file = st.file_uploader("Upload File", type=["xlsx", "docx", "txt", "rtf", "pdf"])
    target = st.selectbox("Target Table", ["subjects", "users"])

    if uploaded_file and st.button("Import Data"):
        try:
            if file_type == "Excel":
                df = pd.read_excel(uploaded_file)
            elif file_type == "Word":
                doc = Document(uploaded_file)
                lines = [p.text for p in doc.paragraphs if p.text.strip()]
                df = pd.DataFrame([line.split(",") for line in lines])
            elif file_type == "Text":
                content = uploaded_file.read().decode("utf-8")
                lines = content.splitlines()
                df = pd.DataFrame([line.split(",") for line in lines if line.strip()])
            elif file_type == "RTF":
                content = uploaded_file.read().decode("utf-8")
                plain = rtf_to_text(content)
                lines = plain.splitlines()
                df = pd.DataFrame([line.split(",") for line in lines if line.strip()])
            elif file_type == "PDF":
                with pdfplumber.open(uploaded_file) as pdf:
                    text = "\n".join([page.extract_text() or "" for page in pdf.pages])
                lines = text.splitlines()
                df = pd.DataFrame([line.split(",") for line in lines if line.strip()])

            st.dataframe(df.head())
            if target == "subjects":
                for _, r in df.iterrows():
                    cursor.execute("""
                        INSERT OR IGNORE INTO subjects (name, description)
                        VALUES (?, ?)
                    """, (r[0], r[1] if len(r) > 1 else ""))
            elif target == "users":
                for _, r in df.iterrows():
                    cursor.execute("""
                        INSERT OR IGNORE INTO users (user_id, first_name, last_name, email, phone_number,
                        gender, date_of_birth, photo, institution, license_key, registered_on)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9], r[10]))
            conn.commit()
            st.success("Data imported successfully.")
        except Exception as e:
            st.error(f"Import error: {e}")

# -----------------------------
# EXPORT SECTION
# -----------------------------
elif menu == "Export Data":
    st.subheader("📅 Export Data")
    export_table = st.selectbox("Select table", ["subjects", "users"])
    export_format = st.selectbox("Format", ["Excel", "PDF"])

    cursor.execute(f"SELECT * FROM {export_table}")
    rows = cursor.fetchall()
    df = pd.DataFrame(rows)

    if st.button("Generate Export"):
        if df.empty:
            st.warning("No data to export.")
        else:
            if export_format == "Excel":
                file = f"{export_table}_export.xlsx"
                df.to_excel(file, index=False)
                with open(file, "rb") as f:
                    st.download_button("Download Excel", f, file_name=file)
            elif export_format == "PDF":
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Arial", size=10)
                col_width = pdf.w / (len(df.columns) + 1)
                for col in df.columns:
                    pdf.cell(col_width, 10, txt=str(col), border=1)
                pdf.ln()
                for _, row in df.iterrows():
                    for val in row:
                        pdf.cell(col_width, 10, txt=str(val), border=1)
                    pdf.ln()
                file = f"{export_table}_export.pdf"
                pdf.output(file)
                with open(file, "rb") as f:
                    st.download_button("Download PDF", f, file_name=file)

# -----------------------------
# DYNAMIC QUERY BUILDER
# -----------------------------
elif menu == "Query Builder":
    st.subheader("🧠 Dynamic Export Query Builder")
    export_table = st.selectbox("Choose table", ["subjects", "users"], key="query_table")
    query_filters = {}

    if export_table == "subjects":
        name = st.text_input("Subject Name")
        desc = st.text_input("Description")
        if name: query_filters['name'] = name
        if desc: query_filters['description'] = desc

    elif export_table == "users":
        uid = st.text_input("User ID")
        email = st.text_input("Email")
        inst = st.text_input("Institution")
        reg_range = st.date_input("Registered On", [datetime.date(2024, 1, 1), datetime.date.today()])
        if uid: query_filters['user_id'] = uid
        if email: query_filters['email'] = email
        if inst: query_filters['institution'] = inst
        query_filters['registered_on_range'] = reg_range

    export_format = st.selectbox("Format", ["Excel", "PDF"], key="query_format")

    if st.button("Run & Export"):
        try:
            query = f"SELECT * FROM {export_table} WHERE 1=1"
            params = []
            for field, val in query_filters.items():
                if field.endswith("_range"):
                    query += " AND DATE(registered_on) BETWEEN ? AND ?"
                    params.extend([val[0].isoformat(), val[1].isoformat()])
                else:
                    query += f" AND {field} LIKE ?"
                    params.append(f"%{val}%")

            cursor.execute(query, params)
            rows = cursor.fetchall()
            df = pd.DataFrame(rows)

            if df.empty:
                st.warning("No matching records.")
            else:
                if export_format == "Excel":
                    file = f"{export_table}_filtered.xlsx"
                    df.to_excel(file, index=False)
                    with open(file, "rb") as f:
                        st.download_button("Download Excel", f, file_name=file)
                elif export_format == "PDF":
                    pdf = FPDF()
                    pdf.add_page()
                    pdf.set_font("Arial", size=10)
                    col_width = pdf.w / (len(df.columns) + 1)
                    for col in df.columns:
                        pdf.cell(col_width, 10, txt=str(col), border=1)
                    pdf.ln()
                    for _, row in df.iterrows():
                        for val in row:
                            pdf.cell(col_width, 10, txt=str(val), border=1)
                        pdf.ln()
                    file = f"{export_table}_filtered.pdf"
                    pdf.output(file)
                    with open(file, "rb") as f:
                        st.download_button("Download PDF", f, file_name=file)

        except Exception as e:
            st.error(f"Query failed: {e}")